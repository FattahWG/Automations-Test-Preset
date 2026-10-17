"""Builds test-cases/LawakStore_TestCases.xlsx from the single source (04 section 4).

Usage:
  python scripts/build_testcases.py                 workbook from the source, with the last known results
  python scripts/build_testcases.py --run <RUN_DIR> also stores the results of that run first

Never edit the workbook by hand: this script overwrites it. A copy of the old file goes to test-cases/backup/.
"""
import argparse
import json
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
TC_DIR = ROOT / "test-cases"
OUT = TC_DIR / "LawakStore_TestCases.xlsx"
LATEST = TC_DIR / "latest-results.json"

HEAD_FILL = PatternFill("solid", fgColor="1D1B16")
HEAD_FONT = Font(bold=True, color="FFFFFF")
FILLS = {"Pass": PatternFill("solid", fgColor="DCFCE7"), "Fail": PatternFill("solid", fgColor="FEE2E2"),
         "Skipped": PatternFill("solid", fgColor="FEF3C7"), "Not run": PatternFill("solid", fgColor="F1F5F9")}
THIN = Border(*(Side(style="thin", color="E2DDD2"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
OWNER = {"UI": "Frontend", "API": "Backend"}

COLUMNS = [("ID", 13), ("Module", 16), ("Sprint", 9), ("Phase", 22), ("Scenario", 18), ("Test Case", 38),
           ("Type", 10), ("Priority", 9), ("Layer", 7), ("Tool", 11), ("Suites", 16), ("Readiness", 12),
           ("Steps", 52), ("Expected Result", 52), ("Result", 9), ("Tested On", 17), ("Run", 26),
           ("Failed Layer", 12), ("Failure Note", 44), ("Owner", 11), ("References", 24)]


def local_time(iso):
    if not iso:
        return ""
    return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone().strftime("%Y-%m-%d %H:%M")


def store_run(run_dir):
    meta = json.loads((run_dir / "run-meta.json").read_text(encoding="utf-8"))
    results = json.loads((run_dir / "case-results.json").read_text(encoding="utf-8"))
    latest = json.loads(LATEST.read_text(encoding="utf-8")) if LATEST.exists() else {}
    for r in results:
        step = next((s for s in r["steps"] if s["status"] == "Fail"), None)
        latest[r["id"]] = {
            "result": r["status"], "testedOn": meta["finishedAt"], "run": f"{meta['scenario']} / {meta['run']}",
            "failedLayer": step["layer"] if step else "", "failureNote": (f"S{step['no']:02d}: {step['actual']}" if step else ""),
            "owner": OWNER.get(step["layer"], "") if step else "",
        }
    LATEST.write_text(json.dumps(latest, indent=2) + "\n", encoding="utf-8")
    return latest


def header(ws, names):
    for i, (name, width) in enumerate(names, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD_FILL, HEAD_FONT, Alignment(vertical="center", wrap_text=True), THIN
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 30


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--run")
    args = ap.parse_args()
    src = json.loads((TC_DIR / "source" / "lawak-store.json").read_text(encoding="utf-8"))
    latest = store_run(Path(args.run).resolve()) if args.run else (json.loads(LATEST.read_text(encoding="utf-8")) if LATEST.exists() else {})
    modules = {m["code"]: m["name"] for m in src["modules"]}
    sprints = {s["id"]: s for s in src["sprints"]}
    phases = {p["id"]: p["name"] for p in src["phases"]}

    if OUT.exists():
        (TC_DIR / "backup").mkdir(exist_ok=True)
        shutil.copy2(OUT, TC_DIR / "backup" / f"LawakStore_TestCases_{datetime.now():%Y%m%d-%H%M%S}.xlsx")

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    # one tab per sprint
    for sp in src["sprints"]:
        ws = wb.create_sheet(sp["name"])
        header(ws, COLUMNS)
        rows = [c for c in src["cases"] if c["sprint"] == sp["id"]]
        for r, c in enumerate(rows, start=2):
            res = latest.get(c["id"], {})
            result = res.get("result", "Not run")
            suites = ", ".join(s.title() for s in c["suites"]) or "Full Cycle only"
            values = [c["id"], modules[c["module"]], sp["name"], phases[sp["phase"]], c["scenario"], c["title"], c["type"],
                      c["priority"], c["layer"], "Bruno CLI" if c["layer"] == "API" else "Playwright", suites,
                      "Waiting for reference" if "[TBD" in json.dumps(c) else "Ready to test",
                      "\n".join(f"{i}. {s['step']}" for i, s in enumerate(c["steps"], 1)),
                      "\n".join(f"{i}. {s['expect']}" for i, s in enumerate(c["steps"], 1)),
                      result, (res.get("testedOn") or "")[:16].replace("T", " "), res.get("run", ""),
                      res.get("failedLayer", ""), res.get("failureNote", ""), res.get("owner", ""), c["refs"]]
            for col, v in enumerate(values, start=1):
                cell = ws.cell(row=r, column=col, value=v)
                cell.alignment, cell.border = WRAP, THIN
            ws.cell(row=r, column=15).fill = FILLS.get(result, FILLS["Not run"])
            ws.cell(row=r, column=15).font = Font(bold=True)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(2, len(rows) + 1)}"
        dv = DataValidation(type="list", formula1='"Pass,Fail,Skipped,Not run"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"O2:O{len(rows) + 1}")

    # summary tab
    summary["A1"] = f"{src['project']} - Test Cases"
    summary["A1"].font = Font(bold=True, size=16, color="C2410C")
    summary["A2"] = f"Generated {datetime.now():%d %b %Y %H:%M} from test-cases/source/lawak-store.json. Do not edit by hand."
    summary["A2"].font = Font(italic=True, color="6B665C")
    head = ["Module"] + [s["name"] for s in src["sprints"]] + ["Total", "Smoke", "Pass", "Fail", "Not run"]
    for i, h in enumerate(head, start=1):
        cell = summary.cell(row=4, column=i, value=h)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, THIN
        summary.column_dimensions[get_column_letter(i)].width = 18 if i == 1 else 11
    row = 5
    for m in src["modules"]:
        cs = [c for c in src["cases"] if c["module"] == m["code"]]
        res = [latest.get(c["id"], {}).get("result", "Not run") for c in cs]
        values = [m["name"]] + [sum(1 for c in cs if c["sprint"] == s["id"]) for s in src["sprints"]] + \
                 [len(cs), sum(1 for c in cs if "smoke" in c["suites"]), res.count("Pass"), res.count("Fail"), res.count("Not run") + res.count("Skipped")]
        for i, v in enumerate(values, start=1):
            summary.cell(row=row, column=i, value=v).border = THIN
        row += 1
    total_row = ["Total"] + [f"=SUM({get_column_letter(i)}5:{get_column_letter(i)}{row - 1})" for i in range(2, len(head) + 1)]
    for i, v in enumerate(total_row, start=1):
        cell = summary.cell(row=row, column=i, value=v)
        cell.font, cell.border = Font(bold=True), THIN

    row += 2
    summary.cell(row=row, column=1, value="Suites").font = Font(bold=True, size=12)
    for s in src["suites"]:
        row += 1
        summary.cell(row=row, column=1, value=s["name"]).font = Font(bold=True)
        summary.cell(row=row, column=2, value=s["description"])
    row += 2
    summary.cell(row=row, column=1, value="Phases and sprints").font = Font(bold=True, size=12)
    for sp in src["sprints"]:
        row += 1
        summary.cell(row=row, column=1, value=sp["name"]).font = Font(bold=True)
        summary.cell(row=row, column=2, value=f"{phases[sp['phase']]}: {sp['goal']}")

    wb.save(OUT)
    print(f"Workbook: {OUT}")


if __name__ == "__main__":
    main()
