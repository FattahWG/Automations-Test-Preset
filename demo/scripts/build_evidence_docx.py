"""Builds Evidence_<scenario>_RunTestN.docx from one run folder.

Usage: python scripts/build_evidence_docx.py --run automation/scenarios/<scenario>/RunTestN
Reads run-meta.json and case-results.json (made by run-suite.js from the Playwright and Bruno reports).
"""
import argparse
import json
from datetime import datetime
from pathlib import Path

from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

from docx_common import BRAND, load_source, mono, new_document, para, picture, table

OWNER = {"UI": "Frontend", "API": "Backend"}
KNOWN_DEFECTS = {
    "TC-CHK-006": {
        "defect": "DEF-001",
        "layer": "API (validation rule)",
        "owner": "Backend",
        "recommendation": "Change the postal code rule from 4-5 digits to exactly 5 digits (regex ^\\d{5}$) in the order "
                          "validation, and add the same check to the checkout form for fast feedback.",
    }
}


def fmt_time(iso):
    return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone().strftime("%d %b %Y %H:%M")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", required=True)
    run_dir = Path(ap.parse_args().run).resolve()
    meta = json.loads((run_dir / "run-meta.json").read_text(encoding="utf-8"))
    results = json.loads((run_dir / "case-results.json").read_text(encoding="utf-8"))
    src = load_source()
    cases = {c["id"]: c for c in src["cases"]}
    modules = {m["code"]: m["name"] for m in src["modules"]}
    sprints = {s["id"]: s["name"] for s in src["sprints"]}

    doc = new_document()
    # ---------- cover ----------
    para(doc, "Lawak Store", bold=True, size=12, color=BRAND)
    doc.add_heading(f"Test Evidence - {meta['scenario']} / {meta['run']}", level=0)
    t = meta["totals"]
    f = meta["filters"]
    filters = ", ".join(f"{k} {v}" for k, v in f.items() if v) or "none"
    table(doc, ["Item", "Value"], [
        ["Application", "Lawak Store (demo shop with admin panel)"],
        ["Suite", meta["suite"].title() + (" (Regression UI + Bruno API)" if meta["suite"] == "full" else "")],
        ["Filters", filters],
        ["Run", f"{meta['run']}  ({fmt_time(meta['startedAt'])} to {fmt_time(meta['finishedAt'])})"],
        ["Target", meta["target"]],
        ["Browser", meta["browser"]],
        ["Command", meta["command"]],
        ["Result", f"{t['pass']} Pass, {t['fail']} Fail, {t['skipped']} Skipped of {t['cases']} cases"],
    ], widths=[1.6, 7.6])
    doc.add_paragraph()
    doc.add_heading("Status of all cases", level=2)
    table(doc, ["ID", "Test case", "Module", "Sprint", "Layer", "Tool", "Result"], [
        [r["id"], r["title"], modules.get(cases.get(r["id"], {}).get("module"), "-"),
         sprints.get(cases.get(r["id"], {}).get("sprint"), "-"), r["layer"], r["tool"], r["status"]]
        for r in results
    ], widths=[1.1, 3.9, 1.4, 0.8, 0.6, 1.0, 0.8], status_col=6)

    # ---------- failure summary (07 section 5) ----------
    failed = [r for r in results if r["status"] == "Fail"]
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Failure summary", level=1)
    if not failed:
        para(doc, "No case failed in this run.")
    else:
        rows = []
        for r in failed:
            step = next((s for s in r["steps"] if s["status"] == "Fail"), None)
            known = KNOWN_DEFECTS.get(r["id"], {})
            layer = known.get("layer") or (step["layer"] if step else r["layer"])
            owner = known.get("owner") or OWNER.get(step["layer"] if step else r["layer"], "QA")
            rows.append([
                r["id"] + (f"\n{known['defect']}" if known else ""), r["title"],
                f"S{step['no']:02d} {step['title']}" if step else "-",
                (step["actual"] if step else r.get("error")) or "-",
                f"{layer} - {owner}",
                known.get("recommendation", "Compare the failed step with the reference. Open the trace to see the page state at the failure."),
                (step or {}).get("evidence") and Path(step["evidence"]).name or "-",
            ])
        table(doc, ["ID", "Test case", "Failed step", "Actual result", "Layer - owner", "Recommendation", "Evidence"], rows,
              widths=[1.0, 1.6, 1.5, 2.0, 1.1, 2.2, 1.2])
        by_owner = {}
        for row in rows:
            by_owner.setdefault(row[4].split(" - ")[-1], []).append(row[0].split("\n")[0])
        doc.add_paragraph()
        for owner, ids in sorted(by_owner.items()):
            para(doc, f"{owner}: {', '.join(ids)}")
    para(doc, f"Check: {t['pass']} Pass + {t['fail']} Fail + {t['skipped']} Skipped = {t['cases']} cases.", italic=True)

    # ---------- per module > scenario > case ----------
    order = [m["code"] for m in src["modules"]]
    results.sort(key=lambda r: (order.index(cases.get(r["id"], {}).get("module", order[0])), r["id"]))
    current_module = None
    current_scenario = None
    for r in results:
        c = cases.get(r["id"], {})
        if c.get("module") != current_module:
            current_module = c.get("module")
            current_scenario = None
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(f"Module: {modules.get(current_module, current_module)}", level=1)
        if c.get("scenario") != current_scenario:
            current_scenario = c.get("scenario")
            doc.add_heading(f"Scenario: {current_scenario}", level=2)
        doc.add_heading(f"{r['id']} - {r['title']}  [{r['status']}]", level=3)
        para(doc, f"Type {c.get('type')} | Priority {c.get('priority')} | {sprints.get(c.get('sprint'))} | "
                  f"Layer {r['layer']} | Tool {r['tool']} | Reference {c.get('refs')}", size=9)
        expects = {i + 1: s["expect"] for i, s in enumerate(c.get("steps", []))}
        steps_src = {i + 1: s["step"] for i, s in enumerate(c.get("steps", []))}
        table(doc, ["No", "Layer", "Step", "Expected result", "Result", "Actual result"], [
            [f"S{s['no']:02d}", s["layer"], steps_src.get(s["no"], s["title"]), expects.get(s["no"], "-"), s["status"], s["actual"]]
            for s in r["steps"]
        ], widths=[0.5, 0.5, 2.6, 2.6, 0.6, 2.4], status_col=4)
        if not r["steps"]:
            para(doc, "No step ran.", italic=True)
        for s in r["steps"]:
            ev = s.get("evidence")
            if not ev:
                continue
            path = run_dir / ev
            if not path.exists():
                continue
            cap = para(doc, f"Evidence {Path(ev).name}", italic=True, size=8)
            if path.suffix.lower() == ".png":
                picture(doc, path)
            elif path.suffix.lower() == ".json":
                text = path.read_text(encoding="utf-8")
                mono(doc, text[:1800] + ("\n..." if len(text) > 1800 else ""))
        if r.get("issue"):
            para(doc, f"Note: {r['issue']}", italic=True, color=RGBColor(0xB9, 0x1C, 0x1C))

    out = run_dir / f"Evidence_{meta['scenario']}_{meta['run']}.docx"
    doc.save(out)
    print(f"Evidence document: {out}")


if __name__ == "__main__":
    main()
